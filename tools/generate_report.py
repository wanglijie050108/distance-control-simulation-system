from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Cm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE = Path(__file__).resolve().parents[1]
TEMPLATE = BASE.parent / "A级实验能力达标线上测试报告模板.docx"
OUTPUT = BASE / "report" / "23009200496_王李杰_计算机科学与技术学院_线上A测报告_待补截图.docx"
ARDUINO_FILE = BASE / "arduino" / "DistanceControl" / "DistanceControl.ino"
PC_FILE = BASE / "pc_host" / "distance_monitor.ps1"


def set_run_font(run, size=12, bold=False, name_cn="宋体", name_en="Times New Roman"):
    run.font.name = name_en
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name_cn)
    run.font.size = Pt(size)
    run.bold = bold


def style_paragraph(paragraph, size=12, bold=False, align=None, line_spacing=1.5):
    if align is not None:
        paragraph.alignment = align
    paragraph.paragraph_format.line_spacing = line_spacing
    paragraph.paragraph_format.space_after = Pt(6)
    for run in paragraph.runs:
        set_run_font(run, size=size, bold=bold)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=16 if level == 1 else 14, bold=True, name_cn="黑体")
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(8)
    return p


def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2)
    else:
        run = p.add_run(text)
        set_run_font(run)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.first_line_indent = Cm(0.74)
    return p


def add_code(doc, title, code):
    add_heading(doc, title, level=2)
    for line in code.splitlines():
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(line if line else " ")
        r.font.name = "Consolas"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
        r.font.size = Pt(8)


def add_screenshot_placeholder(doc, caption, expected):
    p = doc.add_paragraph()
    run = p.add_run(caption)
    set_run_font(run, size=12, bold=True)
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Cm(16)
    cell.height = Cm(9)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_shading(cell, "F2F2F2")
    cp = cell.paragraphs[0]
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cr = cp.add_run(f"【请粘贴仿真截图】\n{expected}")
    set_run_font(cr, size=12, bold=True)
    doc.add_paragraph()


def main():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document(TEMPLATE)

    # 填写封面已有字段。
    replacements = {
        "学院 专业": "学院  计算机科学与技术学院        专业  （请填写）",
        "学号": "学号  23009200496",
        "姓名": "姓名  王李杰",
        "手机 完成日期": "手机  （请填写）        完成日期  2026-09-21",
    }
    for p in doc.paragraphs:
        compact = " ".join(p.text.split())
        for key, value in replacements.items():
            if compact.startswith(key):
                p.text = value
                style_paragraph(p)
                break

    # 从“题目名称”起删除模板占位正文，保留封面和校徽。
    body = doc._element.body
    start = None
    for child in list(body):
        text = "".join(child.itertext())
        if "题目名称" in text:
            start = child
            break
    if start is not None:
        deleting = False
        for child in list(body):
            if child is start:
                deleting = True
            if deleting and child.tag != qn("w:sectPr"):
                body.remove(child)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = title.add_run("题目名称：距离测控仿真系统")
    set_run_font(tr, size=18, bold=True, name_cn="黑体")

    add_heading(doc, "一、题目要求")
    add_body(doc, "使用 Arduino UNO（ATmega328P）搭建 PC 上位机远程距离检测控制系统。Arduino 通过串行接口与 PC 双向通信；PC 发送学生本人学号，Arduino 接收后将当前距离返回 PC，并在 LCD 上显示学号和距离值。")
    add_body(doc, "Arduino 控制直流电机：当实时距离大于个人距离阈值时电机转动；当距离小于或等于阈值时电机停止。LCD 第一行显示“ID：学号”，第二行显示“DIST：距离值”。")
    add_body(doc, "PC 上位机由本人编写，GUI 标题显示学号和姓名，具备串口打开/关闭、学号发送、距离接收显示等功能。串口参数设置为 9600 baud、8 数据位、无校验、1 停止位。")

    add_heading(doc, "二、系统设计")
    add_body(doc, "系统由 Sharp GP2D12 红外测距传感器、Arduino UNO、LM016L 1602 LCD、继电器直流电机驱动电路、COMPIM 串行接口和 PC 上位机组成。传感器输出接 Arduino A0，电机控制端接 D7，LCD 使用 4 位并行方式。")
    add_body(doc, "通信协议采用一行一帧。PC 发送“ID:23009200496”；Arduino 回复“ACK:ID=23009200496”，并周期发送“DIST:距离,MOTOR:ON/OFF”。这种格式便于人工观察，也便于上位机稳定解析。")

    add_heading(doc, "三、距离阈值计算")
    add_body(doc, "学号为 23009200496，末位数字为 6。")
    add_body(doc, "距离阈值 = 30 + 学号末位数 = 30 + 6 = 36 cm。")
    add_body(doc, "因此，距离 > 36 cm 时直流电机转动；距离 ≤ 36 cm 时直流电机停止。")

    add_heading(doc, "四、仿真结果展示")
    add_body(doc, "运行时先配置一对虚拟串口，例如 COM1 ↔ COM2。Proteus 中 COMPIM 使用 COM1，上位机打开 COM2。打开串口后发送完整学号 23009200496，再分别将传感器调到高于和低于阈值的距离并截图。")
    add_screenshot_placeholder(doc, "图 1  距离高于阈值时的运行结果", "建议将传感器调为 45 cm；画面需同时包含学号发送、45.0 cm、LCD、电机转动和上位机接收窗口。")
    add_screenshot_placeholder(doc, "图 2  距离低于或等于阈值时的运行结果", "建议将传感器调为 30 cm；画面需同时包含学号发送、30.0 cm、LCD、电机停止和上位机接收窗口。")

    add_heading(doc, "五、测试结论")
    add_body(doc, "待完成两组 Proteus 联调截图后填写：系统能够完成 PC 与 Arduino 双向串口通信，LCD 与 PC 上位机实时距离显示一致；距离超过 36 cm 时电机启动，距离不超过 36 cm 时电机停止，符合题目要求。")

    add_heading(doc, "六、参考文献")
    refs = [
        "[1] Arduino. Arduino UNO Rev3 Documentation.",
        "[2] Hitachi. HD44780U LCD Controller/Driver Datasheet.",
        "[3] Sharp. GP2D12 Optoelectronic Distance Measuring Sensor Datasheet.",
        "[4] Labcenter Electronics. Proteus Design Suite User Documentation.",
        "[5] Microsoft. Windows PowerShell and .NET System.IO.Ports Documentation.",
    ]
    for ref in refs:
        add_body(doc, ref)

    doc.add_page_break()
    add_heading(doc, "七、程序设计")
    add_code(doc, "7.1 Arduino 程序源代码", ARDUINO_FILE.read_text(encoding="utf-8"))
    doc.add_page_break()
    add_code(doc, "7.2 PC 上位机程序源代码", PC_FILE.read_text(encoding="utf-8"))

    # 统一正文页边距。
    for section in doc.sections:
        section.top_margin = Cm(2.2)
        section.bottom_margin = Cm(2.2)
        section.left_margin = Cm(2.4)
        section.right_margin = Cm(2.4)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
